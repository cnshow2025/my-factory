# -*- coding: utf-8 -*-
"""共用：把貨批與機台資料寫成三分頁的 xlsx（說明／貨批／機台）。"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

F = "Arial"
NAVY = "1E2761"; TEAL = "1C7293"; ICE = "EAF0FA"
GREY = "5A6785"; INK = "16213E"; RED = "C2410C"
_thin = Side(style="thin", color="C9D2E2")
BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

LOT_COLS = [
 ("lot_id",    "批號\n必填。不可重複，直接用現場批號。"),
 ("package",   "封裝／Kit 代號\n必填。同代號＝同一套 Kit；換到不同代號才要換 Kit（60 分）。"),
 ("program",   "測試程式\n必填。package 相同但 program 不同＝只換程式（30 分）。"),
 ("qty",       "顆數\n必填。"),
 ("tpu_sec",   "單 site 單顆測試秒數\n必填。不是整批時間，也不是多 site 的等效時間。"),
 ("arrive_h",  "到廠時間（小時）\n選填。以「現在」為 0。空白＝已經在 WIP 裡。"),
 ("due_h",     "交期（小時）\n選填。以「現在」為 0。空白＝這批不看交期。"),
 ("hot",       "急件\n選填。Y＝急件可插單；空白或 N＝一般批。"),
 ("yield_pct", "良率 %\n選填。填了就用這個值；空白＝用參數的隨機區間。"),
 ("note",      "備註\n選填。不影響模擬。"),
]
MACH_COLS = [
 ("machine_id",     "機台名稱\n必填。"),
 ("sites",          "site 數\n必填。直接當速度倍率：8 site 填 8，同一批只要 1 site 的 1/8 時間。"),
 ("kit_now",        "目前裝的 Kit\n選填。填 package 代號。空白＝空機，第一批要先換 Kit。"),
 ("prog_now",       "目前載入的程式\n選填。空白＝沒有程式，第一批要載入。"),
 ("can_packages",   "能測的封裝\n選填。用直線 | 分隔，例如 QFN48|QFN32。空白＝全部都能測。"),
 ("available_at_h", "幾小時後空出來\n選填。正在跑別的貨或保養中。空白＝現在就能用。"),
]
LOTS_DOC = [
 ("lot_id",    "必填", "批號，不可重複。直接用現場批號。", "LOT-0001"),
 ("package",   "必填", "封裝／Kit 代號。同代號＝同一套 Kit，換到不同代號才要換 Kit（60 分）。", "QFN48"),
 ("program",   "必填", "測試程式名稱或版本。package 相同但 program 不同＝只換程式（30 分）。", "MCU01_V3"),
 ("qty",       "必填", "這批的顆數。", "8000"),
 ("tpu_sec",   "必填", "單一 site 測一顆要幾秒。不是整批時間，也不是多 site 的等效時間。", "2.1"),
 ("arrive_h",  "選填", "到廠時間，以「現在」為 0 的小時數。空白＝已經在 WIP 裡等著。", "0"),
 ("due_h",     "選填", "交期，同樣以現在為 0 的小時數。空白＝這批不看交期。", "36"),
 ("hot",       "選填", "Y＝急件（可插單）。空白或 N＝一般批。", "N"),
 ("yield_pct", "選填", "這個料號的實際良率百分比。空白＝用參數設定的隨機區間。", "96.5"),
 ("note",      "選填", "備註，不影響模擬結果。", "主力機種"),
]
MACH_DOC = [
 ("machine_id",     "必填", "機台名稱。", "T01"),
 ("sites",          "必填", "site 數，直接當速度倍率。8 site 的機台填 8，同一批只要 1 site 的 1/8 時間。", "8"),
 ("kit_now",        "選填", "現在機台上裝著的 Kit，填 package 代號。空白＝空機，第一批要先換 Kit。", "QFN48"),
 ("prog_now",       "選填", "現在載入的測試程式。空白＝沒有程式，第一批要載入。", "MCU01_V3"),
 ("can_packages",   "選填", "這台能測哪些封裝，用直線 | 分隔。空白＝全部都能測。", "QFN48|QFN32"),
 ("available_at_h", "選填", "幾小時後才空出來（正在跑別的貨、保養中）。空白＝現在就能用。", "0"),
]

def _hdr(ws, row, cols, fill=NAVY):
    for c, item in enumerate(cols, start=1):
        key, note = item if isinstance(item, tuple) else (item, None)
        cell = ws.cell(row=row, column=c, value=key)
        cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BOX
        if note:
            cm = Comment(note, "MyFactory"); cm.width = 320; cm.height = 110
            cell.comment = cm

def _put(ws, row, values, blue=False):
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = Font(name=F, size=10, color="0000FF" if blue else INK)
        cell.border = BOX
        cell.alignment = Alignment(horizontal="left" if isinstance(v, str) else "right")

def build(path, lots, machines, note=None, example=False):
    wb = Workbook()

    # ---------- 說明 ----------
    ws = wb.active; ws.title = "說明"
    ws.sheet_view.showGridLines = False
    for col, w in (("A", 18), ("B", 10), ("C", 62), ("D", 30)):
        ws.column_dimensions[col].width = w
    ws["A1"] = "MyFactory 實際貨批模擬 — 填寫範本"
    ws["A1"].font = Font(name=F, bold=True, size=16, color=NAVY)
    ws["A2"] = note or "填好這份檔案後載入模擬器，就會用你真實的 WIP 跑各種排程法則，比較「這批貨全部做完要多久」。"
    ws["A2"].font = Font(name=F, size=10, color=GREY)

    ws["A4"] = "怎麼填"; ws["A4"].font = Font(name=F, bold=True, size=13, color=NAVY)
    how = [
     ("「貨批」分頁", "一列一批待測貨。把現場 WIP 清單貼進來，必填欄位不能空白。"),
     ("「機台」分頁", "一列一台測試機，包含它現在裝著什麼 Kit、載著什麼程式。"),
     ("欄位標題",     "滑鼠移到標題格會跳出這個欄位的中文說明。"),
     ("代號自由命名", "package / program 直接用你現場的名字，種類多少都可以，模擬器會自動吃進去。"),
    ]
    if example:
        how.insert(2, ("藍色字的範例列", "第 2 列起的藍色字是範例，照格式填好自己的資料後把它們刪掉。"))
    for i, (a, b) in enumerate(how):
        r = 5 + i
        ws.cell(row=r, column=1, value=a).font = Font(name=F, bold=True, size=10, color=INK)
        ws.cell(row=r, column=3, value=b).font = Font(name=F, size=10, color=INK)

    ws["A11"] = "貨批 — 欄位說明"; ws["A11"].font = Font(name=F, bold=True, size=13, color=NAVY)
    _hdr(ws, 12, ["欄位", "必填", "說明", "範例"], fill=TEAL)
    for i, row in enumerate(LOTS_DOC):
        r = 13 + i; _put(ws, r, list(row))
        ws.cell(row=r, column=1).font = Font(name=F, bold=True, size=10, color=INK)
        if row[1] == "必填":
            ws.cell(row=r, column=2).font = Font(name=F, bold=True, size=10, color=RED)

    ws["A24"] = "機台 — 欄位說明"; ws["A24"].font = Font(name=F, bold=True, size=13, color=NAVY)
    _hdr(ws, 25, ["欄位", "必填", "說明", "範例"], fill=TEAL)
    for i, row in enumerate(MACH_DOC):
        r = 26 + i; _put(ws, r, list(row))
        ws.cell(row=r, column=1).font = Font(name=F, bold=True, size=10, color=INK)
        if row[1] == "必填":
            ws.cell(row=r, column=2).font = Font(name=F, bold=True, size=10, color=RED)

    ws["A34"] = "填完自動檢查（不用手動算，改資料會跟著變）"
    ws["A34"].font = Font(name=F, bold=True, size=13, color=NAVY)
    _hdr(ws, 35, ["項目", "", "數值", "單位"], fill=TEAL)
    CHK = [
     ("貨批筆數",             "=COUNTA(貨批!A2:A5000)",                        "筆"),
     ("總顆數",               "=SUM(貨批!D2:D5000)",                           "顆"),
     ("純測試工時（1 site）", "=SUMPRODUCT(貨批!D2:D5000,貨批!E2:E5000)/3600", "小時"),
     ("機台台數",             "=COUNTA(機台!A2:A500)",                         "台"),
     ("機隊等效 site 總數",   "=SUM(機台!B2:B500)",                            "site"),
     ("理論最短清空時間",     "=IFERROR(C38/C40,0)",                           "小時（不含改機與當機，排程再好也快不過這個）"),
    ]
    for i, (label, formula, unit) in enumerate(CHK):
        r = 36 + i
        ws.cell(row=r, column=1, value=label).font = Font(name=F, bold=True, size=10, color=INK)
        c = ws.cell(row=r, column=3, value=formula)
        c.font = Font(name=F, size=10, color=INK); c.number_format = "#,##0.0"
        c.fill = PatternFill("solid", fgColor=ICE); c.border = BOX
        ws.cell(row=r, column=4, value=unit).font = Font(name=F, size=10, color=GREY)
    for i, t in enumerate([
        "「理論最短清空時間」＝ 純測試工時 ÷ 機隊等效 site 總數。它假設零改機、零當機、機台永遠不閒著，",
        "所以實際一定比它長。模擬跑出來的數字離它多遠，就是改機與排程吃掉的部分。"]):
        ws.cell(row=43+i, column=1, value=t).font = Font(name=F, size=10, color=GREY)

    # ---------- 貨批 ----------
    ws = wb.create_sheet("貨批")
    _hdr(ws, 1, LOT_COLS)
    for i, r in enumerate(lots):
        _put(ws, 2+i, r, blue=example)
    for i, w in enumerate([13, 12, 13, 9, 13, 12, 11, 7, 10, 26], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"; ws.row_dimensions[1].height = 22
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv); dv.add("H2:H5000")

    # ---------- 機台 ----------
    ws = wb.create_sheet("機台")
    _hdr(ws, 1, MACH_COLS)
    for i, r in enumerate(machines):
        _put(ws, 2+i, r, blue=example)
    for i, w in enumerate([13, 8, 14, 16, 24, 17], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"; ws.row_dimensions[1].height = 22

    wb.save(path)


def add_matrix(path, lots, machines, pkgs):
    """加一頁「能力矩陣」：機台 × 封裝的對照表，外加產能 vs 貨量的檢查。

    這頁是對照用的，模擬器讀的仍然是「機台」分頁的 can_packages。
    """
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter as CL

    wb = load_workbook(path)
    ws = wb.create_sheet("能力矩陣")
    ws.sheet_view.showGridLines = False
    n = len(machines)
    first, last = 3, 2 + len(pkgs)          # 封裝欄的起訖（C 開始）
    r0, r1 = 2, 1 + n                       # 機台資料列的起訖

    head = ["機台", "sites"] + pkgs + ["目前裝著的 Kit", "目前程式", "幾小時後可用"]
    _hdr(ws, 1, head)
    for i, m in enumerate(machines):
        r = 2 + i
        can = m["can"] if m["can"] else pkgs        # 空白＝全部都能測
        row = [m["id"], m["sites"]] + ["✔" if p in can else "" for p in pkgs] \
              + [m["kit"], m["prog"], m["avail"]]
        _put(ws, r, row)
        for c in range(first, last + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = Alignment(horizontal="center")
            if cell.value == "✔":
                cell.font = Font(name=F, size=11, bold=True, color=TEAL)
        # 目前正裝著的那一格反白，一眼看出機隊現在停在什麼配置
        if m["kit"] in pkgs:
            k = ws.cell(row=r, column=first + pkgs.index(m["kit"]))
            k.fill = PatternFill("solid", fgColor="FFF1CC")

    stats = [
        ("能測這種封裝的機台數",   '=COUNTIF({c}%d:{c}%d,"✔")' % (r0, r1)),
        ("能測這種封裝的 site 合計", '=SUMIF({c}%d:{c}%d,"✔",$B$%d:$B$%d)' % (r0, r1, r0, r1)),
        ("目前正裝著這種 Kit 的 site", '=SUMIF($%s$%d:$%s$%d,{c}$1,$B$%d:$B$%d)'
            % (CL(last+1), r0, CL(last+1), r1, r0, r1)),
        ("這批貨的工時（分，1 site）",
            '=SUMPRODUCT((貨批!$B$2:$B$%d={c}$1)*貨批!$D$2:$D$%d*貨批!$E$2:$E$%d)/60'
            % (len(lots)+1, len(lots)+1, len(lots)+1)),
    ]
    base = r1 + 2
    for j, (label, tpl) in enumerate(stats):
        r = base + j
        ws.cell(row=r, column=1, value=label).font = Font(name=F, bold=True, size=10, color=INK)
        for c in range(first, last + 1):
            cell = ws.cell(row=r, column=c, value=tpl.replace("{c}", CL(c)))
            cell.font = Font(name=F, size=10, color=INK)
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="center")
            cell.border = BOX
            cell.fill = PatternFill("solid", fgColor=ICE)
    # 重點列：工時 ÷ 目前裝著的 site ＝ 完全不換 Kit 的話這種封裝要跑多久
    r = base + len(stats)
    ws.cell(row=r, column=1, value="若完全不換 Kit，要跑幾小時").font = Font(name=F, bold=True, size=10, color=RED)
    for c in range(first, last + 1):
        cell = ws.cell(row=r, column=c,
                       value='=IFERROR({c}%d/{c}%d/60,"—")'.replace("{c}", CL(c)) % (base+3, base+2))
        cell.font = Font(name=F, size=10, bold=True, color=INK)
        cell.number_format = "0.0"
        cell.alignment = Alignment(horizontal="center")
        cell.border = BOX
        cell.fill = PatternFill("solid", fgColor="FFE8D9")

    notes = [
        "✔ ＝ 這台機可以測這種封裝（對應「機台」分頁的 can_packages；空白代表該台什麼都能測，這裡會全部打勾）。",
        "黃底 ＝ 這台機「現在」裝著的 Kit，也就是機隊此刻停在什麼配置。",
        "最後一列是重點：把機隊凍結在現在的配置（完全不換 Kit），每種封裝各要跑多久。",
        "哪一種明顯偏高，就是這批貨的瓶頸——代表它的量和現在裝著它的 site 數對不上，得先換 Kit 把配置拉回來。",
        "這頁是對照用的，模擬器讀的仍然是「機台」分頁；改這裡不會改變模擬結果。",
    ]
    for j, t in enumerate(notes):
        ws.cell(row=r + 2 + j, column=1, value=t).font = Font(name=F, size=10, color=GREY)

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 7
    for c in range(first, last + 1):
        ws.column_dimensions[CL(c)].width = max(10, len(pkgs[c - first]) + 3)
    for c in (last + 1, last + 2, last + 3):
        ws.column_dimensions[CL(c)].width = 16
    ws.freeze_panes = "C2"
    ws.row_dimensions[1].height = 22
    wb.save(path)
