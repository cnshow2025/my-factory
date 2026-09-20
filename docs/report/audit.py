from pptx import Presentation
from pptx.util import Emu
import itertools

prs = Presentation("../測試排程分析報告.pptx")
SW, SH = prs.slide_width, prs.slide_height
def inch(v): return v / 914400
print(f"版面 {inch(SW):.2f}in x {inch(SH):.2f}in, 共 {len(prs.slides)} 張\n")

issues = []
for idx, slide in enumerate(prs.slides, 1):
    boxes = []; pics = []
    for sh in slide.shapes:
        if sh.left is None: continue
        l, t = inch(sh.left), inch(sh.top)
        w, h = inch(sh.width or 0), inch(sh.height or 0)
        r, b = l + w, t + h
        name = sh.shape_type
        txt = sh.text_frame.text.strip() if sh.has_text_frame else ""
        # 出界
        if l < -0.01 or t < -0.01 or r > inch(SW)+0.01 or b > inch(SH)+0.01:
            issues.append(f"S{idx} 出界: {txt[:24] or name} → x{l:.2f} y{t:.2f} r{r:.2f} b{b:.2f}")
        # 邊距
        if txt and (l < 0.5 or r > inch(SW)-0.45 or t < 0.25 or b > inch(SH)-0.2):
            issues.append(f"S{idx} 邊距不足: 「{txt[:20]}」 x{l:.2f} r{r:.2f} t{t:.2f} b{b:.2f}")
        # 文字量 vs 容器（粗估：12pt 中文字約 0.17in 寬、行高 0.26in）
        if txt and sh.has_text_frame:
            sizes = [r_.font.size.pt for p in sh.text_frame.paragraphs for r_ in p.runs if r_.font.size]
            fs = max(sizes) if sizes else 12
            cjk = sum(1 for c in txt if ord(c) > 0x2E80)
            ascii_n = len(txt) - cjk
            est_w = cjk*fs/72*1.0 + ascii_n*fs/72*0.5
            lines_needed = max(1, est_w / max(w-0.1, 0.3))
            lines_needed += txt.count("\n")
            need_h = lines_needed * (fs/72*1.45)
            if need_h > h + 0.06:
                issues.append(f"S{idx} 可能溢出: 「{txt[:22]}」 需≈{need_h:.2f}in / 有 {h:.2f}in ({fs:.0f}pt)")
        if txt: boxes.append((l,t,r,b,txt[:18]))
        if sh.shape_type == 13:  # PICTURE
            pics.append((l,t,r,b,"[圖] "+(sh.name or "")))
    # 圖片與文字重疊
    for p in pics:
        for a in boxes:
            ox = min(p[2],a[2]) - max(p[0],a[0]); oy = min(p[3],a[3]) - max(p[1],a[1])
            if ox > 0.05 and oy > 0.05:
                issues.append(f"S{idx} 圖文重疊: {p[4]} ↔ 「{a[4]}」 ({ox:.2f}x{oy:.2f}in)")
    for p, q in itertools.combinations(pics, 2):
        ox = min(p[2],q[2]) - max(p[0],q[0]); oy = min(p[3],q[3]) - max(p[1],q[1])
        if ox > 0.05 and oy > 0.05:
            issues.append(f"S{idx} 圖圖重疊: {p[4]} ↔ {q[4]}")
    # 文字框重疊
    for a, c in itertools.combinations(boxes, 2):
        ox = min(a[2],c[2]) - max(a[0],c[0]); oy = min(a[3],c[3]) - max(a[1],c[1])
        if ox > 0.08 and oy > 0.08:
            issues.append(f"S{idx} 文字重疊: 「{a[4]}」 ↔ 「{c[4]}」 ({ox:.2f}x{oy:.2f}in)")

print(f"發現 {len(issues)} 項:")
for i in issues: print(" -", i)
